# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and contributors
# For license information, please see license.txt
"""The sheets a company fills to bring its chart of accounts and opening balances over.

The chart sheet has the Chart of Accounts Importer's eight columns, then two optional balance
columns. The balances sheet has an account and its balance, or a trial balance's opening and
closing columns, of which the closing ones are used.
"""

import csv
import re
from io import BytesIO, StringIO

import frappe
from frappe import _
from frappe.desk.utils import provide_binary_file
from frappe.utils import cint, cstr, flt
from frappe.utils.csvutils import read_csv_content
from frappe.utils.xlsxutils import make_xlsx, read_xls_file_from_attached_file
from openpyxl import load_workbook

from erpnext.accounts.doctype.chart_of_accounts_importer.chart_of_accounts_importer import (
	build_forest,
	generate_data_from_csv,
	generate_data_from_excel,
	get_file,
	get_root_types,
	get_template,
	validate_columns,
)

IMPORTER_COLUMNS = 8
# the importer's columns a chart sheet must have
CHART_HEADERS = ("Account Name", "Parent Account", "Account Number", "Is Group", "Account Type", "Root Type")
IMPORTER_KEYS = (
	"account_name",
	"parent_account",
	"account_number",
	"parent_account_number",
	"is_group",
	"account_type",
	"root_type",
	"account_currency",
)
# a QuickBooks account type: its top-level account, root type, group and ERPNext account type
ACCOUNT_LIST_TYPES = {
	"bank": ("Assets", "Asset", "Bank Accounts", "Bank"),
	"accounts receivable (a/r)": ("Assets", "Asset", "Accounts Receivable", "Receivable"),
	"other current assets": ("Assets", "Asset", "Other Current Assets", None),
	"fixed assets": ("Assets", "Asset", "Fixed Assets", "Fixed Asset"),
	"other assets": ("Assets", "Asset", "Other Assets", None),
	"accounts payable (a/p)": ("Liabilities", "Liability", "Accounts Payable", "Payable"),
	"credit card": ("Liabilities", "Liability", "Credit Cards", "Bank"),
	"other current liabilities": ("Liabilities", "Liability", "Other Current Liabilities", None),
	"long term liabilities": ("Liabilities", "Liability", "Long Term Liabilities", None),
	"equity": ("Equity", "Equity", None, "Equity"),
	"income": ("Income", "Income", None, "Income Account"),
	"other income": ("Income", "Income", "Other Income", "Income Account"),
	"cost of goods sold": ("Expenses", "Expense", "Cost of Goods Sold", "Cost of Goods Sold"),
	"expenses": ("Expenses", "Expense", None, "Expense Account"),
	"other expense": ("Expenses", "Expense", "Other Expense", "Expense Account"),
}
# Busy's own groups, by their name tidied: the name the chart gives it, the group above, root type
# and account type
BUSY_GROUPS = {
	"assets": ("Assets", None, "Asset", None),
	"liabilities": ("Liabilities", None, "Liability", None),
	"income": ("Income", None, "Income", None),
	"expenses": ("Expenses", None, "Expense", None),
	"capital account": ("Capital Account", None, "Equity", None),
	"reserves surplus": ("Reserves & Surplus", "capital account", "Equity", None),
	"loans liability": ("Loans (Liability)", "liabilities", "Liability", None),
	"secured loans": ("Secured Loans", "loans liability", "Liability", None),
	"unsecured loans": ("Unsecured Loans", "loans liability", "Liability", None),
	"bank o d account": ("Bank O/D Account", "loans liability", "Liability", "Bank"),
	"current liabilities": ("Current Liabilities", "liabilities", "Liability", None),
	"duties taxes": ("Duties & Taxes", "current liabilities", "Liability", "Tax"),
	"provisions expenses payable": ("Provisions/Expenses Payable", "current liabilities", "Liability", None),
	"sundry creditors": ("Sundry Creditors", "current liabilities", "Liability", None),
	"suspense account": ("Suspense Account", "current liabilities", "Liability", None),
	"fixed assets": ("Fixed Assets", "assets", "Asset", None),
	"investments": ("Investments", "assets", "Asset", None),
	"current assets": ("Current Assets", "assets", "Asset", None),
	"bank accounts": ("Bank Accounts", "current assets", "Asset", "Bank"),
	"cash in hand": ("Cash-in-hand", "current assets", "Asset", "Cash"),
	"loans advances asset": ("Loans & Advances (Asset)", "current assets", "Asset", None),
	"securities deposits asset": ("Securities & Deposits (Asset)", "current assets", "Asset", None),
	"sundry debtors": ("Sundry Debtors", "current assets", "Asset", None),
	"stock in hand": ("Stock-in-hand", "current assets", "Asset", "Stock"),
	"sale": ("Sales Accounts", "income", "Income", None),
	"sales": ("Sales Accounts", "income", "Income", None),
	"income direct opr": ("Income (Direct/Opr.)", "income", "Income", None),
	"income indirect": ("Income (Indirect)", "income", "Income", None),
	"purchase": ("Purchase Accounts", "expenses", "Expense", None),
	"purchases": ("Purchase Accounts", "expenses", "Expense", None),
	"expenses direct mfg": ("Expenses (Direct/Mfg.)", "expenses", "Expense", None),
	"expenses indirect admn": ("Expenses (Indirect/Admn.)", "expenses", "Expense", None),
}
BALANCE_COLUMNS = ("Opening Debit", "Opening Credit")
# a ledger named like a tax line, as Indian charts name GST, TDS and TCS ledgers
TAX_WORDS = re.compile(r"\b(gst|cgst|sgst|igst|utgst|tds|tcs)\b", re.IGNORECASE)
# the groups tax ledgers sit under, by their name tidied as get_group_key does
TAX_GROUPS = ("duties taxes", "duties and taxes")
# the charts made at setup have no Equity root, so the importer's rule for all five root types is too strict
REQUIRED_ROOT_TYPES = ("Asset", "Liability", "Income", "Expense")
# customers and suppliers come in as opening invoices, stock as a stock entry, and the temporary
# account takes the difference, so these accounts get no opening balance of their own
NO_OPENING_BALANCE = ("Receivable", "Payable", "Stock", "Temporary")
BALANCE_SHEET_COLUMNS = ("Account Number", "Account Name", "Debit", "Credit")
GROUP_COLUMNS = ("group", "under", "parent_group", "account_group")
NAME_COLUMNS = ("account_name", "account", "particulars", "ledger", "ledger_name", "name")
# how far down a report's headings may be, under its title rows
HEADING_ROWS = 15
# a report's own lines that add up the others
TOTAL_LINE = re.compile(r"^(grand\s+)?total(\s+for\s.*)?:?$", re.IGNORECASE)
# a side written after an amount, like "30,000.00 Cr."
DR_CR = re.compile(r"\b(cr|dr)\.?\s*$", re.IGNORECASE)
# a quantity and the unit Tally writes after it, like "2,450 Kg"
QTY_WITH_UNIT = re.compile(r"^\s*(-?[\d,]*\.?\d+)\s*([^\d\s].*)?$")
# the groups other software keeps customers, suppliers and stock under (Tally, Busy, QuickBooks)
PARTY_GROUPS = {
	"Receivable": (
		"sundry debtors",
		"debtors",
		"accounts receivable",
		"accounts receivable a r",
		"trade receivables",
	),
	"Payable": (
		"sundry creditors",
		"creditors",
		"accounts payable",
		"accounts payable a p",
		"trade payables",
	),
	"Stock": ("stock in hand", "closing stock", "inventory", "inventory asset"),
}


def read_chart_sheet(file_url: str, group_kinds: dict | None = None) -> frappe._dict:
	"""The importer's rows of the sheet, its accounts as a tree, and its opening balances."""
	file_doc, extension = get_file(file_url)
	unknown_groups = []
	rows = get_account_list_rows(file_url)
	if not rows:
		rows, unknown_groups = get_busy_list_rows(file_url, group_kinds) or (None, [])
	if rows:
		dict_rows = [dict(zip(IMPORTER_KEYS, row, strict=False)) for row in rows]
	else:
		dict_rows = read_rows(file_doc, extension, as_dict=True)
		validate_chart_headers(dict_rows)
		rows = [row[:IMPORTER_COLUMNS] for row in read_rows(file_doc, extension)]
	renamed = rename_repeated_accounts(rows, dict_rows)
	set_tax_types(rows)
	validate_columns(rows)
	validate_roots(rows)
	return frappe._dict(
		rows=rows,
		forest=build_forest(rows),
		balances=get_sheet_balances(dict_rows),
		renamed=renamed,
		unknown_groups=unknown_groups,
	)


def set_tax_types(rows: list[list]):
	"""A ledger with no type of its own that is a tax line, like Input CGST or TDS Payable, or any
	ledger under Duties and Taxes, gets the type Tax: tax templates only offer accounts of that type.
	Only on the balance sheet side, so an expense like GST Late Fee stays as it is."""
	by_name = {row[0]: row for row in rows if row and row[0]}

	def get_parents(row):
		seen = set()
		while row and row[1] and row[1] != row[0] and row[0] not in seen:
			seen.add(row[0])
			row = by_name.get(row[1])
			if row:
				yield row

	for row in rows:
		if not row or not row[0] or cint(row[4]) or cstr(row[5]).strip():
			continue
		parents = list(get_parents(row))
		root_type = next((cstr(r[6]).strip() for r in [row, *parents] if cstr(r[6]).strip()), "")
		if root_type not in ("Asset", "Liability"):
			continue
		if TAX_WORDS.search(row[0]) or any(get_group_key(parent[0]) in TAX_GROUPS for parent in parents):
			row[5] = "Tax"


def rename_repeated_accounts(rows: list[list], dict_rows: list[dict]) -> list[str]:
	"""A ledger name used under two parents, like QuickBooks' "Original Cost" under each asset, gets
	its parent's name added: ERPNext keeps one account per name. Returns the new names."""
	counts = {}
	for row in rows:
		if row and row[0] and not cstr(row[2]).strip():
			counts[row[0]] = counts.get(row[0], 0) + 1
	renamed = []
	for row, values in zip(rows, dict_rows, strict=True):
		if not row or counts.get(row[0], 0) < 2 or cstr(row[2]).strip():
			continue
		if cint(row[4]) or row[1] == row[0]:
			frappe.throw(
				_("{0} is in the chart more than once. Give each its own name, or an account number.").format(
					frappe.bold(row[0])
				),
				title=_("Same account twice"),
			)
		row[0] = f"{row[0]} ({row[1]})"
		values["account_name"] = row[0]
		renamed.append(row[0])
	if len(set(renamed)) < len(renamed):
		name = next(name for name in renamed if renamed.count(name) > 1)
		frappe.throw(
			_("{0} is in the chart more than once. Give each its own name, or an account number.").format(
				frappe.bold(name)
			),
			title=_("Same account twice"),
		)
	return renamed


def get_account_list_rows(file_url: str) -> list[list] | None:
	"""A QuickBooks Account List in the importer's columns: each account under a group for its
	QuickBooks type, and a sub-account ("Parent:Child") under its parent. None for any other sheet."""
	report, keys = read_sheet(file_url)
	heading = find_heading(
		keys, lambda row: {"type", "detail_type"} <= set(row) and {"account", "full_name"} & set(row)
	)
	if heading is None:
		return None
	columns = keys[heading]
	name_column = find_column(columns, ("account", "full_name"))
	type_column, detail_column = columns.index("type"), columns.index("detail_type")

	rows, added = [], set()

	def add(name, parent, is_group, account_type, root_type):
		# a ledger name may repeat under other parents: rename_repeated_accounts tells them apart
		key = name if is_group else (name, parent)
		if key not in added:
			added.add(key)
			rows.append([name, parent or name, None, None, is_group, account_type, root_type, None])

	accounts = []
	for row in report[heading + 1 :]:
		name = cstr(get_cell(row.cells, name_column)).strip()
		kind = ACCOUNT_LIST_TYPES.get(cstr(get_cell(row.cells, type_column)).strip().lower())
		if name and kind:
			parts = [part.strip() for part in name.split(":")]
			accounts.append((parts, kind, cstr(get_cell(row.cells, detail_column)).strip().lower()))
	parents = {part for parts, _kind, _detail in accounts for part in parts[:-1]}
	for parts, (root_name, root_type, group, account_type), detail in accounts:
		add(root_name, None, 1, None, root_type)
		under = root_name
		if group:
			add(group, root_name, 1, "Bank" if account_type == "Bank" else None, root_type)
			under = group
		for part in parts[:-1]:
			add(part, under, 1, None, root_type)
			under = part
		if parts[-1] in parents:
			add(parts[-1], under, 1, None, root_type)
		elif "accumulated" in detail:
			add(parts[-1], under, 0, "Accumulated Depreciation", root_type)
		elif detail == "depreciation":
			add(parts[-1], under, 0, "Depreciation", root_type)
		elif "tax payable" in detail and "payroll" not in detail:
			add(parts[-1], under, 0, "Tax", root_type)
		else:
			add(parts[-1], under, 0, account_type, root_type)
	return rows


def get_busy_list_rows(file_url: str, group_kinds: dict | None = None) -> tuple[list[list], list[str]] | None:
	"""A Busy List of Accounts in the importer's columns: each account under its Busy group, placed
	where Busy keeps that group. Customers and suppliers become one Debtors and one Creditors account,
	as they come in as parties. A group Busy doesn't ship, like "Insurance & TPA", is what group_kinds
	says it holds: "Customers" or "Suppliers" make it one receivable or payable ledger of its name,
	else it is a group under Current Assets. Those groups are returned too, to ask about. None for
	any other sheet."""
	group_kinds = group_kinds or {}
	report, keys = read_sheet(file_url)
	heading = find_heading(
		keys,
		lambda row: "group" in row
		and {"name", "account_name"} & set(row)
		and "parent_account" not in row
		and not get_amount_columns([row]),
	)
	if heading is None:
		return None
	columns = keys[heading]
	name_column = find_column(columns, ("name", "account_name"))
	group_column = columns.index("group")
	rows, added, unknown = [], set(), []

	def add(name, parent, is_group, account_type, root_type):
		if name not in added:
			added.add(name)
			rows.append([name, parent or name, None, None, is_group, account_type, root_type, None])

	def add_group(key) -> str:
		name, parent, root_type, account_type = BUSY_GROUPS[key]
		add(_(name), parent and add_group(parent), 1, account_type, root_type)
		return _(name)

	for key in ("assets", "liabilities", "income", "expenses"):
		add_group(key)
	for row in report[heading + 1 :]:
		name = cstr(get_cell(row.cells, name_column)).strip()
		group = cstr(get_cell(row.cells, group_column)).strip()
		if not name or not group or is_total_row(row.cells, name_column):
			continue
		key = get_group_key(group)
		if key not in BUSY_GROUPS:
			if group not in unknown:
				unknown.append(group)
			if group_kinds.get(group) == "Customers":
				add(group, add_group("sundry debtors"), 0, "Receivable", "Asset")
			elif group_kinds.get(group) == "Suppliers":
				add(group, add_group("sundry creditors"), 0, "Payable", "Liability")
			else:
				add(group, add_group("current assets"), 1, None, "Asset")
				add(name, group, 0, None, "Asset")
		elif key == "sundry debtors":
			add(_("Debtors"), add_group(key), 0, "Receivable", "Asset")
		elif key == "sundry creditors":
			add(_("Creditors"), add_group(key), 0, "Payable", "Liability")
		else:
			_name, _parent, root_type, account_type = BUSY_GROUPS[key]
			if key == "fixed assets":
				account_type = "Accumulated Depreciation" if "depreciation" in name.lower() else "Fixed Asset"
			add(name, add_group(key), 0, account_type, root_type)
	return rows, unknown


def validate_chart_headers(rows: list[dict]):
	"""A trial balance or another sheet is not a chart: say so before the importer trips over it."""
	headers = set(rows[0]) if rows else set()
	missing = [label for label in CHART_HEADERS if frappe.scrub(label) not in headers]
	if missing:
		frappe.throw(
			_(
				"This sheet isn't a chart of accounts: it has no {0} column. Download the sample or current chart, fill it in and upload that. Balances from a trial balance go in Opening balances instead."
			).format(", ".join(_(label) for label in missing)),
			title=_("Not a chart of accounts"),
		)


def validate_roots(rows: list[list]):
	# the importer gives a top-level account itself as its parent
	roots = [row for row in rows if row[0] and row[1] == row[0]]
	for row in roots:
		if row[6] not in get_root_types():
			frappe.throw(
				_(
					"{0} is a top-level account, so its Root Type must be Asset, Liability, Equity, Income or Expense."
				).format(frappe.bold(row[0])),
				title=_("Wrong root type"),
			)

	missing = set(REQUIRED_ROOT_TYPES) - {row[6] for row in roots}
	if missing:
		frappe.throw(
			_("Add a top-level account for: {0}").format(", ".join(sorted(missing))),
			title=_("Top-level account missing"),
		)


def read_rows(file_doc, extension: str, as_dict: bool = False) -> list:
	if extension == "csv":
		return generate_data_from_csv(file_doc, as_dict=as_dict)
	return generate_data_from_excel(file_doc, extension, as_dict=as_dict)


def read_balance_sheet(file_url: str) -> list[frappe._dict]:
	"""The lines of a balances sheet or trial balance, as exported from other software: under a few
	title rows, with group lines that total the lines indented under them, and a total at the end.
	Each line knows the groups it is under; group lines are marked, totals left out."""
	rows, keys = read_sheet(file_url)
	heading = find_heading(keys, get_amount_columns, above=3)
	amounts = balance_column = None
	if heading is None:
		# one balance column, its side in the amount: "25,00,000.00 Cr" or a minus
		heading = find_heading(keys, lambda row: get_balance_column(row) is not None)
		if heading is not None:
			balance_column = get_balance_column(keys[heading])
	else:
		amounts = get_amount_columns(keys[max(heading - 3, 0) : heading + 1])
	if heading is None:
		frappe.throw(
			_(
				"The sheet needs an Account Name column and Debit and Credit columns (or Closing Debit and Closing Credit), or one Closing Balance column with Dr and Cr."
			),
			title=_("Columns not found"),
		)
	debit_column, credit_column = amounts or (None, None)
	# the name heading may sit rows above the amounts, like Tally's "Particulars"
	name_column = next(
		(row.index(key) for row in reversed(keys[: heading + 1]) for key in NAME_COLUMNS if key in row), 0
	)
	number_column = find_column(keys[heading], ("account_number",))
	# a report may name each account's group in a column, as Busy does, instead of indenting
	group_column = find_column(keys[heading], GROUP_COLUMNS)

	lines = []
	for row in rows[heading + 1 :]:
		name = cstr(get_cell(row.cells, name_column)).strip()
		if not name or is_total_row(row.cells, name_column):
			continue
		number = cstr(get_cell(row.cells, number_column)).strip()
		if balance_column is None:
			debit = get_amount(get_cell(row.cells, debit_column))
			credit = get_amount(get_cell(row.cells, credit_column))
		else:
			balance = get_signed_amount(get_cell(row.cells, balance_column))
			debit, credit = max(balance, 0), max(-balance, 0)
		# QuickBooks names a sub-account "Parent:Child"
		*sheet_parents, leaf = (part.strip() for part in name.split(":"))
		if not leaf or not all(sheet_parents):
			sheet_parents, leaf = [], name
		lines.append(
			frappe._dict(
				account_name=leaf,
				account_number=number,
				label=f"{number} - {name}" if number else name,
				debit=debit,
				credit=credit,
				indent=get_cell(row.indents, name_column) or 0,
				group=cstr(get_cell(row.cells, group_column)).strip(),
				parent=sheet_parents[-1] if sheet_parents else None,
				sheet_parents=sheet_parents,
			)
		)

	balances, parents = [], []
	for line, after in zip(lines, [*lines[1:], None], strict=True):
		while parents and parents[-1].indent >= line.indent:
			parents.pop()
		line.groups = (
			[parent.account_name for parent in parents]
			+ ([line.group] if line.group else [])
			+ line.pop("sheet_parents")
		)
		# a line with lines indented under it is a group, and its amount only totals theirs
		line.is_group = bool(after and after.indent > line.indent)
		if line.is_group:
			parents.append(line)
		if line.is_group or line.debit or line.credit:
			balances.append(line)
	return balances


def read_report(file_url: str) -> list[frappe._dict]:
	"""Every row of an uploaded csv or Excel sheet as it is, with how far each cell's text is
	indented: reports from accounting software indent each account under its group. A csv keeps no
	indent, as Frappe's reader trims its values."""
	file_doc, extension = get_file(file_url)
	content = file_doc.get_content()
	if extension == "xlsx":
		sheet = load_workbook(BytesIO(content), data_only=True).active
		return [
			frappe._dict(cells=[cell.value for cell in row], indents=[get_indent(cell) for cell in row])
			for row in sheet.iter_rows()
		]
	rows = (
		read_csv_content(content, use_sniffer=True)
		if extension == "csv"
		else read_xls_file_from_attached_file(content)
	)
	return [frappe._dict(cells=row, indents=[get_indent(value) for value in row]) for row in rows]


def get_indent(cell) -> float:
	value = getattr(cell, "value", cell)
	if not isinstance(value, str):
		return 0
	indent = cell.alignment.indent if hasattr(cell, "alignment") else 0
	return indent + len(value) - len(value.lstrip())


def get_amount_columns(heading_rows: list[list[str]]) -> tuple[int, int] | None:
	"""The debit and credit columns of the last heading row, the closing ones first: they are the
	balances on the day. Any heading with debit or Dr in it counts, like Busy's "Dr. Amount (Rs.)",
	and a label like "Closing Balance" in a row above counts for the columns under it."""
	above = get_labels_above(heading_rows, max(len(row) for row in heading_rows))
	found = {}
	for index, key in enumerate(heading_rows[-1]):
		words = set(key.split("_"))
		debit, credit = bool(words & {"debit", "dr"}), bool(words & {"credit", "cr"})
		# a "Dr/Cr" column says which side, it holds no amount
		if debit != credit:
			rank = 0 if "closing" in above[index] else 2 if "opening" in above[index] else 1
			found.setdefault(("debit" if debit else "credit", rank), index)
	return next(
		(
			(found[("debit", rank)], found[("credit", rank)])
			for rank in (0, 1, 2)
			if ("debit", rank) in found and ("credit", rank) in found
		),
		None,
	)


def get_amount(value) -> float:
	"""A number cell, or the number in a text cell past a currency like "Rs." or "$"; brackets or a
	minus before it make it negative. The decimal mark is the last of "." and ",", so "1.200,50" is
	1200.5; a lone "," before three digits groups thousands, like "1,20,000"."""
	if not isinstance(value, str):
		return flt(value)
	number = re.search(r"\d[\d.,]*", value)
	if not number:
		return 0
	text = number.group().rstrip(".,")
	if "," in text and "." in text:
		decimal = "," if text.rfind(",") > text.rfind(".") else "."
	elif "," in text:
		decimal = None if len(text.rsplit(",", 1)[1]) == 3 else ","
	else:
		decimal = None if text.count(".") > 1 else "."
	other = "." if decimal == "," else ","
	text = text.replace(other, "")
	text = text.replace(",", ".") if decimal == "," else text.replace(".", "") if decimal is None else text
	amount = flt(text)
	return -amount if value.strip().startswith("(") or "-" in value[: number.start()] else amount


def get_balance_column(keys: list[str]) -> int | None:
	"""The one closing balance column of a heading row that names the accounts too."""
	if not set(keys) & set(NAME_COLUMNS):
		return None
	return next(
		(
			index
			for index, key in enumerate(keys)
			if {"balance", "closing"} & set(key.split("_")) and "opening" not in key.split("_")
		),
		None,
	)


def get_signed_amount(value) -> float:
	"""A balance as debit minus credit: "1,200.00 Cr" and -1200 are credits."""
	amount = get_amount(value)
	side = get_dr_cr(value)
	if side:
		return -abs(amount) if side == "cr" else abs(amount)
	return amount


def get_dr_cr(value) -> str | None:
	"""The side a cell names, "dr" or "cr": a Dr/Cr cell, or an amount ending in one, like "30,000.00 Cr."."""
	match = DR_CR.search(value.strip()) if isinstance(value, str) else None
	return match.group(1).lower() if match else None


def get_qty(value) -> tuple[float, str | None]:
	"""A quantity cell and the unit written in it, like 2,450 and "Kg" from "2,450 Kg"."""
	if not isinstance(value, str):
		return flt(value), None
	match = QTY_WITH_UNIT.match(value)
	if not match:
		return 0, None
	return flt(match.group(1).replace(",", "")), (match.group(2) or "").strip() or None


def read_sheet(file_url: str, get_key=None) -> tuple[list[frappe._dict], list[list[str]]]:
	"""The rows of a sheet as read_report reads them, and the cells of each as column keys."""
	rows = read_report(file_url)
	get_key = get_key or get_column_key
	return rows, [[get_key(value) for value in row.cells] for row in rows]


def find_heading(keys: list[list[str]], is_heading, above: int = 0) -> int | None:
	"""The first row among the title rows that is_heading takes for the headings. With above, it is
	given the row with up to that many rows over it, for headings written across rows."""
	return next(
		(
			i
			for i in range(min(len(keys), HEADING_ROWS))
			if is_heading(keys[max(i - above, 0) : i + 1] if above else keys[i])
		),
		None,
	)


def find_column(headers: list[str], names: tuple[str, ...]) -> int | None:
	"""The column with the first of these names."""
	return next((headers.index(name) for name in names if name in headers), None)


def get_labels_above(heading_rows: list[list[str]], width: int) -> list[set[str]]:
	"""The words over each column in these heading rows: a merged label is in its first cell only,
	and covers the empty ones after it."""
	above = [set() for _ in range(width)]
	for row in heading_rows:
		label = ""
		for index in range(width):
			label = get_cell(row, index) or label
			above[index] |= set(label.split("_"))
	return above


def is_total_row(cells: list, column: int | None = None) -> bool:
	"""A report's own line that adds up the others: any text cell, or the one in column, says so."""
	values = cells if column is None else [get_cell(cells, column)]
	return any(TOTAL_LINE.match(value.strip()) for value in values if isinstance(value, str))


def get_column_key(value) -> str:
	return re.sub(r"[^a-z0-9]+", "_", cstr(value).lower()).strip("_")


def get_group_key(name) -> str:
	"""A group's name as BUSY_GROUPS and PARTY_GROUPS keep it: "Duties & Taxes" is "duties taxes"."""
	return get_column_key(name).replace("_", " ")


def get_cell(cells: list, index: int | None):
	return cells[index] if index is not None and index < len(cells) else None


def get_party_group_type(names: list[str]) -> str | None:
	"""Receivable, Payable or Stock when one of these names is the group other software keeps
	customers, suppliers or stock under."""
	keys = {get_group_key(name) for name in names}
	return next((account_type for account_type, groups in PARTY_GROUPS.items() if keys & set(groups)), None)


def download_sheet(file_name: str, rows: list[list], sheet_name: str):
	"""Send these rows to the browser as an Excel file."""
	provide_binary_file(file_name, "xlsx", make_xlsx(rows, sheet_name).getvalue())


@frappe.whitelist()
def download_opening_balances(company: str):
	"""The accounts that can take an opening balance, with the draft balances filled in."""
	onboarding = frappe.get_doc("Company Setup", company)
	onboarding.check_permission("read")
	drafts = get_draft_balances(onboarding)
	accounts = frappe.get_list(
		"Account",
		filters={
			"company": company,
			"is_group": 0,
			"disabled": 0,
			"account_type": ("not in", NO_OPENING_BALANCE),
		},
		fields=["name", "account_number", "account_name"],
		order_by="lft",
	)
	rows = [list(BALANCE_SHEET_COLUMNS)]
	for account in accounts:
		debit, credit = drafts.get(account.name, (0, 0))
		rows.append([account.account_number, account.account_name, debit or None, credit or None])

	download_sheet(
		f"{frappe.scrub(company)}_opening_balances",
		rows,
		"Opening Balances",
	)


def get_draft_balances(onboarding) -> dict[str, tuple[float, float]]:
	"""The draft debit and credit of each account."""
	balances = {}
	for row in onboarding.opening_balances:
		if row.account and row.status != "Posted":
			debit, credit = balances.get(row.account, (0, 0))
			balances[row.account] = (debit + flt(row.debit), credit + flt(row.credit))
	return balances


def get_sheet_balances(rows: list[dict]) -> list[frappe._dict]:
	"""The rows with an opening debit or credit. Only accounts without sub-accounts can have one."""
	parents = {
		(cstr(row.get("parent_account_number")).strip(), cstr(row.get("parent_account")).strip())
		for row in rows
	}
	balances = []
	for row in rows:
		debit, credit = flt(row.get("opening_debit")), flt(row.get("opening_credit"))
		if not debit and not credit:
			continue

		number, name = cstr(row.get("account_number")).strip(), cstr(row.get("account_name")).strip()
		label = f"{number} - {name}" if number else name
		if cint(row.get("is_group")) or (number, name) in parents:
			frappe.throw(
				_(
					"{0} has sub-accounts, so it can't have an opening balance. Put it on one of its sub-accounts."
				).format(frappe.bold(label)),
				title=_("Balance on a group account"),
			)
		balances.append(
			frappe._dict(account_name=name, account_number=number, label=label, debit=debit, credit=credit)
		)
	return balances


@frappe.whitelist()
def download_chart_template(company: str, template: str):
	"""An Excel sheet to fill: a sample chart, or the company's chart with its draft opening balances."""
	onboarding = frappe.get_doc("Company Setup", company)
	onboarding.check_permission("read")

	rows = get_sample_rows(company)
	if template == "Current Chart":
		rows = [rows[0], *get_current_chart_rows(onboarding)]

	download_sheet(
		(
			f"{frappe.scrub(company)}_chart_of_accounts"
			if template == "Current Chart"
			else "sample_chart_of_accounts"
		),
		rows,
		"Chart of Accounts",
	)


def get_sample_rows(company: str) -> list[list]:
	"""The importer's sample chart, with the balance columns added."""
	rows = list(csv.reader(StringIO(get_template("Sample Template", company).getvalue())))
	return [rows[0] + list(BALANCE_COLUMNS), *rows[1:]]


def get_current_chart_rows(onboarding) -> list[list]:
	"""The company's accounts in the sheet's columns, with their draft opening balances."""
	accounts = frappe.get_list(
		"Account",
		filters={"company": onboarding.company},
		fields=[
			"name",
			"account_name",
			"account_number",
			"parent_account",
			"is_group",
			"account_type",
			"root_type",
			"account_currency",
		],
		order_by="lft",
	)
	by_name = {account.name: account for account in accounts}
	balances = get_draft_balances(onboarding)

	rows = []
	for account in accounts:
		parent = by_name.get(account.parent_account) or frappe._dict()
		debit, credit = balances.get(account.name, (0, 0))
		rows.append(
			[
				account.account_name,
				parent.account_name,
				account.account_number,
				parent.account_number,
				account.is_group,
				account.account_type,
				account.root_type,
				account.account_currency,
				debit or None,
				credit or None,
			]
		)
	return rows


def strip_title_rows(doctype: str, file_url: str) -> str:
	"""The sheet as Data Import reads it, as a new file: without the title rows other software puts
	above its headings (the first row of more than one cell is the headings), its Total lines, and with
	a heading naming the doctype, like QuickBooks' "Customer", read as its name field. The same file
	when there is nothing to change."""
	rows = [list(row.cells) for row in read_report(file_url)]
	start = find_heading(rows, lambda row: sum(1 for value in row if cstr(value).strip()) > 1) or 0
	heading = rows[start] if rows else []
	title_field = frappe.get_meta(doctype).get_title_field()
	names = {get_column_key(doctype), "vendor" if doctype == "Supplier" else None} - {None}
	renamed = [
		frappe.get_meta(doctype).get_label(title_field)
		if isinstance(value, str) and get_column_key(value) in names and title_field != "name"
		else value
		for value in heading
	]
	body = [row for row in rows[start + 1 :] if not is_total_row(row)]
	if not start and renamed == heading and len(body) == len(rows) - 1:
		return file_url
	rows = [renamed, *body]
	content = make_xlsx(rows, _("Sheet1")).getvalue()
	return (
		frappe.get_doc(
			{
				"doctype": "File",
				"file_name": f"{frappe.scrub(doctype)}_import.xlsx",
				"content": content,
				"is_private": 1,
			}
		)
		.insert()
		.file_url
	)
