# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and Contributors
# See license.txt

from contextlib import contextmanager
from io import BytesIO
from unittest.mock import patch

import frappe
from frappe.utils import flt
from openpyxl import Workbook
from openpyxl.styles import Alignment

from erpnext.setup.doctype.company_setup import company_setup_steps
from erpnext.setup.doctype.company_setup.company_setup import create_company_setup
from erpnext.setup.doctype.company_setup.company_setup_steps import HOOK_NAME, get_step_paths
from erpnext.tests.utils import ERPNextTestSuite

COMPANY = "_Test Company"
NEW_COMPANY = "_Test Onboarding Company"

# a small chart from an old system, with the balances it closed with
SHEET = [
	"Account Name,Parent Account,Account Number,Parent Account Number,Is Group,Account Type,Root Type,Account Currency,Opening Debit,Opening Credit",
	"Assets,,1000,,1,,Asset,INR,,",
	"Current Assets,Assets,1100,1000,1,,Asset,INR,,",
	"Cash in Hand,Current Assets,1110,1100,0,Cash,Asset,INR,25000,",
	"HDFC Bank,Current Assets,1120,1100,0,Bank,Asset,INR,180000,",
	"Sundry Debtors,Current Assets,1130,1100,0,Receivable,Asset,INR,95000,",
	"Fixed Assets,Assets,1200,1000,1,,Asset,INR,,",
	"Furniture,Fixed Assets,1210,1200,0,Fixed Asset,Asset,INR,200000,",
	"Liabilities,,2000,,1,,Liability,INR,,",
	"Sundry Creditors,Liabilities,2100,2000,0,Payable,Liability,INR,,82000",
	"Owner Capital,Liabilities,2200,2000,0,,Liability,INR,,418000",
	"Income,,4000,,1,,Income,INR,,",
	"Sales,Income,4100,4000,0,Income Account,Income,INR,,",
	"Expenses,,5000,,1,,Expense,INR,,",
	"Purchases,Expenses,5100,5000,0,Cost of Goods Sold,Expense,INR,,",
	"Rent,Expenses,5200,5000,0,,Expense,INR,,",
]


def get_onboarding():
	create_company_setup(COMPANY)
	doc = frappe.get_doc("Company Setup", COMPANY)
	# as when the form opens, so steps added since the record was made are there
	doc.check_steps()
	return doc


def make_company() -> str:
	return (
		frappe.get_doc(
			{
				"doctype": "Company",
				"company_name": NEW_COMPANY,
				"abbr": "_TOC",
				"default_currency": "INR",
				"country": "India",
				"chart_of_accounts": "Standard",
			}
		)
		.insert()
		.name
	)


def make_user(email: str, role: str) -> str:
	if not frappe.db.exists("User", email):
		frappe.get_doc(
			{"doctype": "User", "email": email, "first_name": role, "send_welcome_email": 0}
		).insert()
	frappe.get_doc("User", email).add_roles(role)
	return email


def make_sheet(rows: list[str], balances: bool = True) -> str:
	"""Upload the sheet as a CSV file. Without balances, the two balance columns are left empty."""
	if not balances:
		rows = [rows[0], *(",".join([*row.split(",")[:8], "", ""]) for row in rows[1:])]
	return (
		frappe.get_doc(
			{"doctype": "File", "file_name": "chart.csv", "content": "\n".join(rows), "is_private": 1}
		)
		.insert()
		.file_url
	)


def step_status(doc) -> dict:
	return {row.step_key.split(".")[-1]: row.status for row in doc.steps}


def extra_step(onboarding) -> dict:
	"""A step another app could add through the hook."""
	return {"title": "Extra step", "description": "", "done": False}


def broken_step(onboarding) -> dict:
	"""A step of another app with a bug in it."""
	raise ZeroDivisionError


@contextmanager
def hooked_steps(steps_by_app: dict[str, list]):
	"""Fake the step hook of these apps; the others keep their real hook. New apps come last."""
	real_get_hooks = frappe.get_hooks
	installed = frappe.get_installed_apps()
	apps = [*installed, *(app for app in steps_by_app if app not in installed)]

	def get_hooks(hook=None, *args, **kwargs):
		if hook == HOOK_NAME and kwargs.get("app_name") in steps_by_app:
			return list(steps_by_app[kwargs["app_name"]])
		return real_get_hooks(hook, *args, **kwargs)

	with (
		patch.object(frappe, "get_installed_apps", return_value=apps),
		patch.object(frappe, "get_hooks", side_effect=get_hooks),
	):
		yield


def step_names(**steps_by_app) -> list[str]:
	"""The order of the steps, by name, with these apps' hooks."""
	with hooked_steps(steps_by_app):
		return [path.rsplit(".", 1)[-1] for path in get_step_paths()]


# ERPNext's own steps before Go Live, in order
ERPNEXT_STEPS = [
	"migration_date",
	"chart_of_accounts",
	"taxes",
	"opening_balances",
	"bank_accounts",
	"customers",
	"opening_sales_invoices",
	"open_sales_orders",
	"suppliers",
	"opening_purchase_invoices",
	"open_purchase_orders",
	"items",
	"warehouses",
	"opening_stock",
	"fixed_assets",
	"boms",
	"open_work_orders",
	"job_workers",
	"open_subcontracting_orders",
	"check_books",
]
# steps whose status comes only from this company's onboarding, not from shared records
OWN_STEPS = ("migration_date", "chart_of_accounts", "opening_balances", "check_books", "go_live")


def steps_with(after: dict | None = None, before: dict | None = None) -> list[str]:
	"""ERPNext's steps with other steps put after or before some of them."""
	steps = []
	for step in [*ERPNEXT_STEPS, "go_live"]:
		steps += (before or {}).get(step, [])
		steps.append(step)
		steps += (after or {}).get(step, [])
	return steps


def step_key(doc, name: str) -> str:
	return next(row.step_key for row in doc.steps if row.step_key.endswith(f".{name}"))


class TestCompanySetup(ERPNextTestSuite):
	def test_onboarding_is_made_renamed_and_deleted_with_the_company(self):
		company = make_company()
		doc = frappe.get_doc("Company Setup", company)
		doc.check_steps()
		self.assertEqual(doc.status, "In Progress")
		self.assertEqual(
			step_status(doc),
			{
				step: "Not Started"
				for step in (
					"migration_date",
					"chart_of_accounts",
					"opening_balances",
					"check_books",
					"go_live",
				)
			}
			| {step: status for step, status in step_status(doc).items() if step not in OWN_STEPS},
		)
		# every step shows, with or without a bank ledger or a trial balance
		self.assertEqual(list(step_status(doc)), [*ERPNEXT_STEPS, "go_live"])

		for name in OWN_STEPS:
			self.assertRaises(frappe.ValidationError, doc.skip_step, step_key(doc, name))
			self.assertEqual(step_status(doc)[name], "Not Started")
		key = step_key(doc, "opening_sales_invoices")
		doc.skip_step(key)
		# a skipped step stays skipped when the steps are checked again
		doc.check_steps()
		self.assertEqual(step_status(doc)["opening_sales_invoices"], "Skipped")
		doc.skip_step(key, skip=0)
		self.assertEqual(step_status(doc)["opening_sales_invoices"], "Not Started")
		# a done step stays done
		make_customer("_Test Setup Skip Customer")
		doc.check_steps()
		doc.skip_step(step_key(doc, "customers"))
		self.assertEqual(step_status(doc)["customers"], "Done")

		renamed = "_Test Renamed Onboarding Company"
		frappe.rename_doc("Company", company, renamed, force=True)
		self.assertEqual(frappe.db.get_value("Company Setup", renamed, "company"), renamed)

		frappe.delete_doc("Company", renamed)
		self.assertFalse(frappe.db.exists("Company Setup", renamed))

	def test_a_hooked_step_shows_and_a_failing_one_does_not_stop_the_setup(self):
		extra = f"{__name__}.extra_step"
		with hooked_steps({"app_one": [{"step": extra, "before": ERPNEXT_STEPS[2]}]}):
			doc = get_onboarding()
			doc.save()
		self.assertEqual(doc.steps[2].step_key, extra)
		self.assertEqual(doc.steps[2].step, "Extra step")

		# a step no app adds any more is dropped
		doc.save()
		self.assertNotIn(extra, [row.step_key for row in doc.steps])

		broken = f"{__name__}.broken_step"
		with hooked_steps({"app_one": [broken]}):
			doc = get_onboarding()
			doc.save()
			company = make_company()
		self.assertNotIn(broken, [row.step_key for row in doc.steps])
		self.assertTrue(frappe.db.exists("Company Setup", company))

	def test_check_books_needs_temporary_opening_at_zero_and_nothing_left_to_post(self):
		doc = self.posting_setup()
		cash = get_account(self.company, account_type="Cash")
		# a balance missing from the sheet: Temporary Opening takes the 500
		doc.save_opening_balances([{"account": cash, "debit": 500}])
		doc.post_opening_balances()
		self.assertEqual(step_status(get_setup(self.company))["check_books"], "Not Started")

		doc.save_opening_balances([{"account": self.capital, "credit": 500}])
		doc.post_opening_balances()
		self.assertEqual(step_status(get_setup(self.company))["check_books"], "Done")

		doc.reload()
		doc.save_opening_invoices("Customer", [{"party_name": "X", "outstanding_amount": 10}])
		doc.save_opening_invoices(
			"Supplier",
			[{"party_name": "Y", "outstanding_amount": 10}, {"party_name": "Z", "outstanding_amount": 10}],
		)
		doc = get_setup(self.company)
		self.assertEqual(step_status(doc)["check_books"], "Not Started")
		self.assertEqual(
			company_setup_steps.check_books(doc)["summary"], "1 invoice and 2 bills still to post"
		)

	def posting_setup(self):
		"""A new company's setup, worked on as an Accounts Manager, migrating on 01-04-2026."""
		self.company = make_company()
		self.capital = get_account(self.company, root_type="Equity", account_type=("!=", "Temporary"))
		frappe.set_user(make_user("onboarding.accounts@example.com", "Accounts Manager"))
		self.addCleanup(frappe.set_user, "Administrator")
		doc = frappe.get_doc("Company Setup", self.company)
		doc.migration_date = "2026-04-01"
		doc.save()
		return doc

	def test_go_live(self):
		doc = get_onboarding()
		# nothing is changed while invoices post in the background
		with patch("erpnext.setup.doctype.company_setup.company_setup.is_job_enqueued", return_value=True):
			self.assertRaises(
				frappe.ValidationError,
				doc.save_opening_invoices,
				"Customer",
				[{"party_name": "Onboarding Busy Customer", "outstanding_amount": 100}],
			)
			self.assertFalse(doc.check_steps())
		doc.go_live()
		self.assertEqual(doc.status, "Live")
		self.assertEqual(step_status(doc)["go_live"], "Done")

		# with money in Temporary Opening, only on a confirmation
		doc = frappe.get_doc("Company Setup", make_company())
		doc.migration_date = "2026-04-01"
		customer = frappe.get_all("Customer", pluck="customer_name", limit=1)[0]
		doc.save_opening_invoices("Customer", [{"party_name": customer, "outstanding_amount": 500}])
		doc.create_opening_invoices("Customer")
		doc.reload()

		self.assertRaises(frappe.ValidationError, doc.go_live)
		doc.reload()
		self.assertEqual(doc.status, "In Progress")

		doc.go_live(accept_temporary_opening=True)
		self.assertEqual(doc.status, "Live")
		doc.run_method("onload")
		summary = doc.get_onload().steps[step_key(doc, "go_live")]["summary"]
		self.assertIn("still in Temporary Opening", summary)


class TestStepOrder(ERPNextTestSuite):
	"""Where steps from other apps go."""

	def test_steps_go_before_or_in_place_of_the_step_they_name(self):
		self.assertEqual(
			step_names(
				app_one=[
					{"step": "app_one.steps.gst", "before": "taxes"},
					{"step": "app_one.steps.bank", "before": "check_books"},
				]
			),
			steps_with(before={"taxes": ["gst"], "check_books": ["bank"]}),
		)

		# at the same place, in install order
		self.assertEqual(
			step_names(
				app_one=[{"step": "app_one.steps.gst", "before": "check_books"}],
				app_two=[{"step": "app_two.steps.employees", "before": "check_books"}],
			),
			steps_with(before={"check_books": ["gst", "employees"]}),
		)

		# a step of an app installed before can be named too
		self.assertEqual(
			step_names(
				app_one=["app_one.steps.employees"],
				app_two=[{"step": "app_two.steps.payroll", "before": "employees"}],
			),
			[*ERPNEXT_STEPS, "payroll", "employees", "go_live"],
		)

		# a step takes the place of the step it replaces
		self.assertEqual(
			step_names(app_one=[{"step": "app_one.steps.gst", "replaces": "opening_balances"}]),
			[step if step != "opening_balances" else "gst" for step in [*ERPNEXT_STEPS, "go_live"]],
		)
		# with the same name, and replacing a step of the same app
		with hooked_steps(
			{
				"app_one": [{"step": "app_one.steps.customers", "replaces": "customers"}],
				"app_two": [
					{"step": "app_two.steps.employees", "before": "suppliers"},
					{"step": "app_two.steps.payroll", "replaces": "employees"},
				],
			}
		):
			paths = get_step_paths()
		self.assertEqual(
			[path.rsplit(".", 1)[-1] for path in paths], steps_with(before={"suppliers": ["payroll"]})
		)
		self.assertIn("app_one.steps.customers", paths)
		# only the first step of that name is left out, not another app's step named the same
		with hooked_steps(
			{
				"app_one": [{"step": "app_one.steps.gst", "replaces": "opening_balances"}],
				"app_two": ["app_two.steps.opening_balances"],
			}
		):
			paths = get_step_paths()
		self.assertIn("app_two.steps.opening_balances", paths)
		self.assertNotIn("erpnext.setup.doctype.company_setup.company_setup_steps.opening_balances", paths)
		# Go Live is never replaced
		self.assertEqual(
			step_names(app_one=[{"step": "app_one.steps.gst", "replaces": "go_live"}])[-2:],
			["gst", "go_live"],
		)


def get_account(company: str, **filters) -> str:
	return frappe.get_all(
		"Account", filters={"company": company, "is_group": 0, **filters}, pluck="name", limit=1
	)[0]


def make_customer(name: str, disabled: int = 0) -> str:
	user = frappe.session.user
	frappe.set_user("Administrator")
	customer = frappe.get_doc(
		{"doctype": "Customer", "customer_name": name, "customer_type": "Company", "disabled": disabled}
	).insert()
	frappe.set_user(user)
	return customer.name


def set_party_account(doctype: str, name: str, company: str, account: str):
	user = frappe.session.user
	frappe.set_user("Administrator")
	doc = frappe.get_doc(doctype, name)
	doc.set("accounts", [{"company": company, "account": account}])
	doc.save()
	frappe.set_user(user)


def has_year(date: str) -> bool:
	return bool(
		frappe.db.exists(
			"Fiscal Year", {"year_start_date": ("<=", date), "year_end_date": (">=", date), "disabled": 0}
		)
	)


def make_xlsx_file(name: str, rows: list[tuple], indent_column: int | None = None) -> str:
	"""An Excel file of these rows; with indent_column, a row's extra last value is that cell's indent."""
	workbook = Workbook()
	width = max(len(row) for row in rows) - (1 if indent_column else 0)
	for row in rows:
		workbook.active.append(row[:width])
		if indent_column and len(row) > width:
			workbook.active.cell(workbook.active.max_row, indent_column).alignment = Alignment(indent=row[-1])
	content = BytesIO()
	workbook.save(content)
	return (
		frappe.get_doc({"doctype": "File", "file_name": name, "content": content.getvalue(), "is_private": 1})
		.insert()
		.file_url
	)


def make_file(name: str, rows: list[str]) -> str:
	return (
		frappe.get_doc({"doctype": "File", "file_name": name, "content": "\n".join(rows), "is_private": 1})
		.insert()
		.file_url
	)


def get_balance(account: str, voucher_no: str | None = None) -> float:
	filters = {"account": account, "is_cancelled": 0}
	if voucher_no:
		filters["voucher_no"] = voucher_no
	rows = frappe.get_all("GL Entry", filters=filters, fields=["debit", "credit"])
	return flt(sum(flt(row.debit) - flt(row.credit) for row in rows), 2)


def make_stock_item(code: str, uom: str, conversions: dict | None = None) -> str:
	if not frappe.db.exists("Item", code):
		frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": code,
				"item_name": code,
				"item_group": "All Item Groups",
				"stock_uom": uom,
				"is_stock_item": 1,
				"uoms": [{"uom": uom, "conversion_factor": 1}]
				+ [
					{"uom": other, "conversion_factor": factor}
					for other, factor in (conversions or {}).items()
				],
			}
		).insert()
	return code


def get_setup(company: str):
	doc = frappe.get_doc("Company Setup", company)
	doc.check_steps()
	return doc


def make_bom(item: str, raw_material: str, company: str) -> str:
	bom = frappe.get_doc(
		{
			"doctype": "BOM",
			"item": item,
			"company": company,
			"quantity": 1,
			"items": [{"item_code": raw_material, "qty": 2, "rate": 10}],
		}
	).insert()
	bom.submit()
	return bom.name
