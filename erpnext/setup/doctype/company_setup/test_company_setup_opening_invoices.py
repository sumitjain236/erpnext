# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and Contributors
# See license.txt

from contextlib import contextmanager
from io import BytesIO

import frappe
from frappe.custom.doctype.property_setter.property_setter import make_property_setter
from openpyxl import load_workbook

from erpnext.accounts.party import get_party_account
from erpnext.accounts.utils import get_fiscal_year
from erpnext.setup.doctype.company_setup import company_setup_steps
from erpnext.setup.doctype.company_setup.opening_invoices import download_new_parties, get_invoice_accounts
from erpnext.setup.doctype.company_setup.test_company_setup import (
	COMPANY,
	get_account,
	make_company,
	make_customer,
	make_file,
	make_user,
	step_key,
	step_status,
)
from erpnext.tests.utils import ERPNextTestSuite


@contextmanager
def day_first_dates():
	"""Dates written day first, as Indian sites usually do."""
	frappe.local.user_date_format = "dd-mm-yyyy"
	try:
		yield
	finally:
		frappe.local.user_date_format = None


class TestOpeningInvoices(ERPNextTestSuite):
	"""Unpaid customer and supplier invoices on the Selling and Buying tabs."""

	def setUp(self):
		self.enterContext(self.change_settings("Accounts Settings", {"delete_linked_ledger_entries": 1}))
		user = make_user("onboarding.accounts@example.com", "Accounts Manager")
		for role in ("Sales Manager", "Purchase Manager"):
			make_user(user, role)
		self.company = make_company()
		frappe.set_user(user)
		self.doc = frappe.get_doc("Company Setup", self.company)
		self.customer = frappe.get_all("Customer", fields=["name", "customer_name"], limit=1)[0]

	def tearDown(self):
		frappe.set_user("Administrator")
		super().tearDown()

	def post_two_sources(self):
		"""One invoice from a sheet and one typed in, posted; their invoice names by source."""
		self.doc.migration_date = "2026-04-01"
		row = {"party_name": self.customer.customer_name, "invoice_date": "2026-03-01"}
		self.doc.save_opening_invoices(
			"Customer", [row | {"invoice_number": "OB-SHEET-1", "outstanding_amount": 100}], "delhi.xlsx"
		)
		self.doc.save_opening_invoices(
			"Customer",
			[
				*[row.as_dict() for row in self.doc.opening_sales_invoices],
				row | {"invoice_number": "OB-TYPED-1", "outstanding_amount": 50},
			],
		)
		self.assertEqual([row.source for row in self.doc.opening_sales_invoices], ["delhi.xlsx", None])
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()
		return {row.source or "": row.invoice for row in self.doc.opening_sales_invoices}

	def test_undo_posting_takes_back_one_source_only(self):
		invoices = self.post_two_sources()
		posted = self.doc.get_posted_entries("Customer")
		self.assertEqual(
			sorted((group["source"], group["count"], group["amount"]) for group in posted["groups"]),
			[("", 1, 50), ("delhi.xlsx", 1, 100)],
		)

		# a row of the same sheet that isn't posted goes with it
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": self.customer.customer_name,
					"invoice_number": "OB-SHEET-2",
					"outstanding_amount": 70,
					"source": "delhi.xlsx",
				}
			],
		)
		result = self.doc.undo_posting("Customer", "delhi.xlsx")
		self.doc.reload()
		self.assertEqual((result.removed, result.unposted, result.kept), (1, 1, 0))
		self.assertFalse(frappe.db.exists("Sales Invoice", invoices["delhi.xlsx"]))
		self.assertEqual(frappe.db.get_value("Sales Invoice", invoices[""], "docstatus"), 1)
		self.assertEqual([row.invoice_number for row in self.doc.opening_sales_invoices], ["OB-TYPED-1"])
		# the undone invoice is kept as an undone row, with its account, to download and upload again
		self.assertEqual(
			[(row.invoice_number, row.account, row.source) for row in self.doc.undone_sales_invoices],
			[("OB-SHEET-1", get_party_account("Customer", self.customer.name, self.company), "delhi.xlsx")],
		)
		self.assertEqual(
			[
				(group["source"], group["count"])
				for group in self.doc.get_posted_entries("Customer")["undone"]
			],
			[("delhi.xlsx", 1)],
		)

		# the right sheet posts again under the same number
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": self.customer.customer_name,
					"invoice_number": "OB-SHEET-1",
					"invoice_date": "2026-03-01",
					"outstanding_amount": 120,
				}
			],
			"delhi-fixed.xlsx",
		)
		self.doc.create_opening_invoices("Customer")
		self.assertEqual(frappe.db.get_value("Sales Invoice", "OB-SHEET-1", "docstatus"), 1)
		self.doc.reload()
		self.assertEqual(self.doc.undone_sales_invoices, [])

	def test_undo_posting_of_picked_invoices_keeps_the_rest_of_the_sheet(self):
		self.doc.migration_date = "2026-04-01"
		row = {"party_name": self.customer.customer_name, "invoice_date": "2026-03-01"}
		self.doc.save_opening_invoices(
			"Customer",
			[
				row | {"invoice_number": "OB-PICK-1", "outstanding_amount": 100},
				row | {"invoice_number": "OB-PICK-2", "outstanding_amount": 60},
			],
			"delhi.xlsx",
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		result = self.doc.undo_posting("Customer", "delhi.xlsx", ["OB-PICK-2"])
		self.doc.reload()
		self.assertEqual((result.removed, result.unposted), (1, 0))
		self.assertFalse(frappe.db.exists("Sales Invoice", "OB-PICK-2"))
		self.assertEqual(
			[(row.invoice_number, row.status) for row in self.doc.opening_sales_invoices],
			[("OB-PICK-1", "Created")],
		)
		self.assertEqual([row.invoice_number for row in self.doc.undone_sales_invoices], ["OB-PICK-2"])

		# a form save from a page opened before the undo can't bring back or drop undone rows
		self.doc.flags.posting = False
		self.doc.undone_sales_invoices = []
		self.doc.save()
		self.doc.reload()
		self.assertEqual([row.invoice_number for row in self.doc.undone_sales_invoices], ["OB-PICK-2"])

	def test_undo_posting_keeps_an_invoice_with_a_payment_and_is_refused_once_live(self):
		from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry

		invoices = self.post_two_sources()
		payment = get_payment_entry("Sales Invoice", invoices["delhi.xlsx"], party_amount=40)
		payment.update({"reference_no": "OB-PAY-1", "reference_date": "2026-04-02"})
		payment.insert()
		payment.submit()

		result = self.doc.undo_posting("Customer", "delhi.xlsx")
		self.assertEqual((result.removed, result.kept), (0, 1))
		self.assertEqual(frappe.db.get_value("Sales Invoice", invoices["delhi.xlsx"], "docstatus"), 1)

		self.doc.db_set("status", "Live")
		self.assertRaises(frappe.ValidationError, self.doc.undo_posting, "Customer", "delhi.xlsx")

	def test_a_closed_year_stops_posting_and_undo_with_a_plain_reason(self):
		invoices = self.post_two_sources()
		frappe.get_doc(
			{
				"doctype": "Period Closing Voucher",
				"name": "_Test Setup PCV",
				"company": self.company,
				"period_end_date": "2026-03-31",
				"docstatus": 1,
			}
		).db_insert()
		closed = "Last year is closed in ERPNext, so opening entries can't"
		with self.assertRaisesRegex(frappe.ValidationError, closed + " be undone"):
			self.doc.undo_posting("Customer", "delhi.xlsx")
		self.assertEqual(frappe.db.get_value("Sales Invoice", invoices["delhi.xlsx"], "docstatus"), 1)

		self.doc.save_opening_invoices(
			"Customer",
			[
				*[row.as_dict() for row in self.doc.opening_sales_invoices],
				{"party_name": self.customer.customer_name, "outstanding_amount": 70},
			],
		)
		with self.assertRaisesRegex(frappe.ValidationError, closed + " be posted"):
			self.doc.create_opening_invoices("Customer")

	def test_opening_invoices_made_outside_are_listed_apart(self):
		self.post_two_sources()
		tool = frappe.new_doc("Opening Invoice Creation Tool")
		tool.update({"company": self.company, "invoice_type": "Sales"})
		tool.append(
			"invoices",
			{
				"party": self.customer.name,
				"posting_date": "2026-03-31",
				"due_date": "2026-03-31",
				"outstanding_amount": 70,
				"temporary_opening_account": get_account(self.company, account_type="Temporary"),
			},
		)
		tool.make_invoices()

		outside = self.doc.get_posted_entries("Customer")["outside"]
		self.assertEqual((outside["count"], outside["amount"]), (1, 70))
		self.assertEqual(outside["rows"][0].party, self.customer.name)

	def test_save_links_customers_by_name_or_as_picked(self):
		self.doc.save_opening_invoices(
			"Customer",
			[
				{"party_name": f"  {self.customer.customer_name.upper()} ", "outstanding_amount": 100},
				{"party_name": "Onboarding New Customer", "outstanding_amount": 50},
			],
		)
		rows = self.doc.opening_sales_invoices
		self.assertEqual([row.party for row in rows], [self.customer.name, None])
		self.assertEqual(self.doc.get_invoice_names("Customer")["new"], ["Onboarding New Customer"])
		self.assertEqual(step_status(self.doc)["opening_sales_invoices"], "Not Started")
		self.doc.run_method("onload")
		step = self.doc.get_onload().steps[step_key(self.doc, "opening_sales_invoices")]
		self.assertTrue(step["started"])

		self.doc.save_opening_invoices(
			"Customer",
			[{"party_name": "Sheet Name Only", "party": self.customer.name, "outstanding_amount": 100}],
		)
		row = self.doc.opening_sales_invoices[0]
		self.assertEqual((row.party_name, row.party), ("Sheet Name Only", self.customer.name))
		self.assertEqual(self.doc.get_invoice_names("Customer")["new"], [])

	def test_names_almost_like_a_customer_or_each_other_are_pointed_out(self):
		with self.subTest("one letter off a customer"):
			almost = self.customer.customer_name[:-1] + "x"
			self.doc.save_opening_invoices("Customer", [{"party_name": almost, "outstanding_amount": 10}])
			self.assertEqual(
				self.doc.get_invoice_names("Customer")["suggestions"], {almost: self.customer.name}
			)

		with self.subTest("only a typo, not another name sharing a word"):
			homes = make_customer("Patil Homes Pvt Ltd")
			make_customer("Joshi Developers")
			self.doc.save_opening_invoices(
				"Customer",
				[
					{"party_name": "Pátil Homs", "outstanding_amount": 10},
					{"party_name": "Joshi Builders", "outstanding_amount": 20},
				],
			)
			self.assertEqual(self.doc.get_invoice_names("Customer")["suggestions"], {"Pátil Homs": homes})

		with self.subTest("Indian name styles"):
			icici = make_customer("ICICI Bank Ltd")
			godrej = make_customer("Godrej & Boyce Mfg Co Ltd")
			adani = make_customer("Adani Realty Pvt Ltd")
			mahindra = make_customer("Mahindra Lifespace Developers Ltd")
			self.doc.save_opening_invoices(
				"Customer",
				[
					{"party_name": "M/s I.C.I.C.I. BANK", "outstanding_amount": 10},
					{"party_name": "Godrej and Boyce Mfg. Co. Ltd.", "outstanding_amount": 20},
					{"party_name": "M/s ADANI REALTY", "outstanding_amount": 30},
					{"party_name": "M/s MAHINDRA LIFESPACE", "outstanding_amount": 40},
				],
			)
			self.assertEqual(
				[row.party for row in self.doc.opening_sales_invoices], [icici, godrej, adani, None]
			)
			self.assertEqual(
				self.doc.get_invoice_names("Customer")["suggestions"], {"M/s MAHINDRA LIFESPACE": mahindra}
			)

		with self.subTest("two spellings of one new name"):
			self.doc.save_opening_invoices(
				"Customer",
				[
					{"party_name": "Ramesh Kumar Agarwal", "outstanding_amount": 10},
					{"party_name": "Ramesh Kumar Agrawal", "outstanding_amount": 20},
					{"party_name": "Global Hospitals Hyderabad", "outstanding_amount": 30},
				],
			)
			self.assertEqual(
				self.doc.get_invoice_names("Customer")["look_alikes"],
				[["Ramesh Kumar Agarwal", "Ramesh Kumar Agrawal"]],
			)

	def test_a_similar_name_links_only_when_one_party_matches(self):
		cafe = make_customer("Café Brew Lane Pvt. Ltd.")
		twin = make_customer("Twin Traders")
		make_customer("Twin Traders Ltd")
		self.doc.save_opening_invoices(
			"Customer",
			[
				{"party_name": "cafe brew-lane", "outstanding_amount": 100},
				{"party_name": "TWIN TRADERS LLP", "outstanding_amount": 50},
				{"party_name": "Twin Traders", "outstanding_amount": 70},
			],
		)
		self.assertEqual([row.party for row in self.doc.opening_sales_invoices], [cafe, None, twin])
		names = self.doc.get_invoice_names("Customer")
		self.assertEqual(names["similar"], {"cafe brew-lane": cafe})
		self.assertEqual(names["new"], ["TWIN TRADERS LLP"])

	def test_a_disabled_or_foreign_currency_customer_is_pointed_out_and_not_posted(self):
		disabled = make_customer("Onboarding Disabled Customer", disabled=1)
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer", [{"party_name": "Onboarding Disabled Customer", "outstanding_amount": 100}]
		)
		self.assertEqual(self.doc.get_invoice_names("Customer")["disabled"], [disabled])
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		row = self.doc.opening_sales_invoices[0]
		self.assertEqual((row.party, row.status, row.error), (disabled, "Draft", None))
		self.assertEqual(frappe.db.count("Customer", {"customer_name": "Onboarding Disabled Customer"}), 1)

		frappe.set_user("Administrator")
		receivable = get_account(self.company, account_type="Receivable")
		export = frappe.get_doc(
			{
				"doctype": "Account",
				"account_name": "Debtors - Export",
				"company": self.company,
				"parent_account": frappe.db.get_value("Account", receivable, "parent_account"),
				"account_type": "Receivable",
				"account_currency": "USD",
			}
		).insert()
		customer = frappe.get_doc("Customer", make_customer("Onboarding Export Customer"))
		customer.default_currency = "USD"
		customer.append("accounts", {"company": self.company, "account": export.name})
		customer.save()
		frappe.set_user("onboarding.accounts@example.com")
		self.doc.save_opening_invoices(
			"Customer", [{"party_name": "Onboarding Export Customer", "outstanding_amount": 100}]
		)
		names = self.doc.get_invoice_names("Customer")
		self.assertEqual(names["other_currency"], {customer.name: "USD"})
		# its USD account is no account to pick, so saving the rows as shown doesn't stop on it
		self.add_receivable_account("Debtors - Domestic")
		self.assertNotIn(customer.name, self.doc.get_invoice_names("Customer")["party_accounts"])
		# nor does the books' preview count it
		self.assertEqual(self.doc.preview_opening_balances()["left_out_invoices"], 1)
		# not even on a ledger in the company's currency, as the Selling card leaves it out
		row = self.doc.opening_sales_invoices[0]
		frappe.db.set_value(row.doctype, row.name, "account", receivable)
		self.doc.reload()
		self.assertEqual(self.doc.preview_opening_balances()["left_out"]["currencies"], {"USD": 1})
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		row = self.doc.opening_sales_invoices[0]
		self.assertEqual((row.status, row.error), ("Draft", None))

	def test_invoice_after_the_migration_date_is_refused(self):
		self.doc.migration_date = "2026-04-01"
		self.assertRaises(
			frappe.ValidationError,
			self.doc.save_opening_invoices,
			"Customer",
			[{"party_name": "X", "invoice_date": "2026-04-05", "outstanding_amount": 10}],
		)

	def test_bills_sheets_from_tally_and_excel(self):
		name = self.customer.customer_name
		with self.subTest("a receivables report"):
			result = self.doc.import_opening_invoices(
				"Customer",
				make_file(
					"receivables.csv",
					[
						"Customer,Invoice No,Invoice Date,Due Date,Outstanding",
						f"{name},OB-INV-1,2026-02-01,2026-03-03,1200",
						"Onboarding New Customer,OB-INV-2,,,800",
						f"{name},,,,-300",
					],
				),
			)
			self.assertEqual(
				result,
				{
					"added": 3,
					"advances": 1,
					"skipped": 0,
					"repeated": 0,
					"unused_columns": [],
					"bad_dates": [],
					"unknown_ledgers": [],
					"post_to": None,
					"source": "receivables.csv",
				},
			)
			self.assertEqual(
				[(row.invoice_number, row.outstanding_amount) for row in self.doc.opening_sales_invoices],
				[("OB-INV-1", 1200), ("OB-INV-2", 800), (None, -300)],
			)

		with self.subTest("a tracker whose Ledger column names the account"):
			_, export = self.add_receivable_account("Debtors - Export")
			result = self.doc.import_opening_invoices(
				"Customer",
				make_file(
					"tracker.csv",
					[
						"Customer Name,Ledger,Invoice No,Outstanding",
						f"{name},Debtors - Export,OB-LDG-1,700",
						f"{name},Old Ledger,OB-LDG-2,300",
					],
				),
			)
			self.assertEqual(
				[(row.invoice_number, row.account) for row in self.doc.opening_sales_invoices],
				[("OB-LDG-1", export), ("OB-LDG-2", None)],
			)
			self.assertEqual(result["unknown_ledgers"], ["Old Ledger"])

		with self.subTest("Tally's bills report, as exported"), day_first_dates():
			report = [
				"Kapoor Interiors",
				"Office 7, Baner Road, Pune",
				"Bills Receivable",
				"1-Apr-2025 to 31-Mar-2026",
				"Date,Ref. No.,Party's Name,Pending Amount,Due on,Overdue by days",
				f'18-01-2026,KI/141,{name},"2,36,000.00 Dr",17-02-2026,42',
				f'25-03-2026,On Account,{name},"30,000.00 Cr",,',
				f",On Account,{name},12500,,",
				"Total,,,,,",
			]
			result = self.doc.import_opening_invoices("Customer", make_file("bills.csv", report))
			self.assertEqual(
				result,
				{
					"added": 3,
					"advances": 1,
					"skipped": 0,
					"repeated": 0,
					"unused_columns": ["Overdue by days"],
					"bad_dates": [],
					"unknown_ledgers": [],
					"post_to": None,
					"source": "bills.csv",
				},
			)
			self.assertEqual(
				[(row.invoice_number, row.outstanding_amount) for row in self.doc.opening_sales_invoices],
				[("KI/141", 236000), (None, -30000), (None, 12500)],
			)
			self.assertEqual(str(self.doc.opening_sales_invoices[0].invoice_date), "2026-01-18")

			self.doc.import_opening_invoices(
				"Supplier",
				make_file(
					"payables.csv",
					["Bills Payable", "Party's Name,Pending Amount", 'Sai Marble,"20,000.00 Dr"'],
				),
			)
			self.assertEqual(self.doc.opening_purchase_invoices[0].outstanding_amount, -20000)

		with self.subTest("Tally's bills with Dr and Cr columns"):
			report = [
				"Bills Receivable",
				"Date,Ref. No.,Party's Name,Pending Amount,,Due on",
				",,,Dr,Cr,",
				f"2026-01-10,KI/1,{name},5000,,2026-02-09",
				f"2026-03-01,On Account,{name},,800,",
				",,Grand Total,5000,800,",
			]
			result = self.doc.import_opening_invoices("Customer", make_file("tally_drcr.csv", report))
			self.assertEqual(
				[(row.invoice_number, row.outstanding_amount) for row in self.doc.opening_sales_invoices],
				[("KI/1", 5000), (None, -800)],
			)
			self.assertEqual(result["advances"], 1)

		with self.subTest("a total in another column is not an invoice"):
			report = [
				"Debtors outstanding - 31 Mar 2026",
				"Customer,Bill No,Bill Date,Amount (Rs),Their Bill No",
				f"{name},FM/311,2026-01-31,212400,PO-9",
				",,Total,212400,",
			]
			result = self.doc.import_opening_invoices("Customer", make_file("excel.csv", report))
			self.assertEqual(
				[(row.invoice_number, row.outstanding_amount) for row in self.doc.opening_sales_invoices],
				[("FM/311", 212400)],
			)
			self.assertEqual(result["unused_columns"], ["Their Bill No"])

		with self.subTest("amounts with a currency and a name that starts with Total"):
			report = [
				"Customer,Invoice No,Outstanding",
				'Total Environment Systems,T-1,"Rs. 1,250.00"',
				"Total,,1250",
			]
			self.doc.import_opening_invoices("Customer", make_file("rs.csv", report))
			self.assertEqual(
				[(row.party_name, row.outstanding_amount) for row in self.doc.opening_sales_invoices],
				[("Total Environment Systems", 1250)],
			)

	def test_bills_sheets_from_busy_and_quickbooks(self):
		name = self.customer.customer_name
		with self.subTest("Busy, with a Dr/Cr column and a named total"), day_first_dates():
			report = [
				"Bill-wise Outstanding (Receivable) as on 31-03-2026",
				"Party Name,Bill No.,Bill Date,Bill Amt.,Pending Amt.,Dr/Cr,Their Ref",
				f"{name},SI/1012,12-02-2026,145600,98750,Dr,PO-7",
				f"{name},,22-03-2026,,50000,Cr,",
				"Total,,,,48750,Dr,",
			]
			result = self.doc.import_opening_invoices("Customer", make_file("busy.csv", report))
			self.assertEqual(
				[(row.invoice_number, row.outstanding_amount) for row in self.doc.opening_sales_invoices],
				[("SI/1012", 98750), (None, -50000)],
			)
			self.assertEqual(result["unused_columns"], ["Bill Amt.", "Their Ref"])

		with self.subTest("Busy's bill-wise report, grouped by party"):
			report = [
				"Outstanding Analysis (Bill-wise) - Sundry Debtors",
				"Party / Bill No.,Bill Date,Bill Amt.,Pending Amt.,Dr/Cr",
				f"M/s {name},,,,",
				'BLR-1306,25/11/2025,"2,23,430.00","2,23,430.00",Dr',
				'BLR-1336,07/02/2026,"2,57,912.00","77,373.60",Dr',
				'Total,,,"3,00,803.60",Dr',
				'Grand Total,,,"3,00,803.60",Dr',
			]
			self.doc.import_opening_invoices("Customer", make_file("busy_billwise.csv", report))
			self.assertEqual(
				[
					(row.party_name, row.invoice_number, row.outstanding_amount)
					for row in self.doc.opening_sales_invoices
				],
				[(f"M/s {name}", "BLR-1306", 223430), (f"M/s {name}", "BLR-1336", 77373.6)],
			)

		with self.subTest("Busy's columns with day-first dates"), day_first_dates():
			self.doc.import_opening_invoices(
				"Customer",
				make_file(
					"busy.csv",
					[
						"Party Name,Bill No.,Bill Date,Due Date,Pending Amount",
						f"{name},B-1,10-05-2026,09-06-2026,500",
					],
				),
			)
			row = self.doc.opening_sales_invoices[0]
			self.assertEqual((str(row.invoice_date), str(row.due_date)), ("2026-05-10", "2026-06-09"))
			self.assertEqual((row.invoice_number, row.outstanding_amount), ("B-1", 500))

		with self.subTest("QuickBooks open invoices, grouped by customer"):
			report = [
				"Maple Street Goods",
				"Open Invoices",
				",Date,Transaction Type,Num,Due Date,Amount,Open Balance",
				f"{name},,,,,,",
				',2026-05-12,Invoice,1042,2026-06-11,"1,250.00","1,250.00"',
				",2026-06-20,Invoice,1058,2026-07-20,845.50,600.00",
				f'Total for {name},,,,,,"$1,850.00"',
				"Quechee Gorge Village Inc.,,,,,,",
				",2026-06-18,Payment,,,,(1500.00)",
			]
			self.doc.import_opening_invoices("Customer", make_file("qbo.csv", report))
			self.assertEqual(
				[
					(row.party_name, row.invoice_number, row.outstanding_amount)
					for row in self.doc.opening_sales_invoices
				],
				[
					(name, "1042", 1250),
					(name, "1058", 600),
					("Quechee Gorge Village Inc.", None, -1500),
				],
			)

	def test_create_posts_known_customers_and_never_adds_one(self):
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": self.customer.customer_name,
					"invoice_number": "OB-SINV-0001",
					"due_date": "2026-01-01",
					"outstanding_amount": 1200,
				},
				{"party_name": "Onboarding New Customer", "outstanding_amount": 800},
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		rows = self.doc.opening_sales_invoices
		self.assertEqual([row.status for row in rows], ["Created", "Failed"], [row.error for row in rows])
		self.assertEqual(rows[1].error, "Onboarding New Customer is not in ERPNext yet.")
		invoice = frappe.get_doc("Sales Invoice", "OB-SINV-0001")
		self.assertEqual(
			(invoice.is_opening, invoice.docstatus, invoice.outstanding_amount), ("Yes", 1, 1200)
		)
		# no date, and no year in the books for one: dated the day before the start
		self.assertEqual(str(invoice.posting_date), "2026-03-31")
		self.assertEqual(str(invoice.due_date), "2026-03-31")
		self.assertFalse(frappe.db.exists("Customer", {"customer_name": "Onboarding New Customer"}))
		self.assertEqual(company_setup_steps.get_temporary_opening_balance(self.company), -1200)

	def test_a_missing_default_receivable_account_is_named_before_posting(self):
		# a chart whose only receivable ledger is in another currency leaves the default empty
		frappe.db.set_value("Company", self.company, "default_receivable_account", None)
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer", [{"party_name": self.customer.customer_name, "outstanding_amount": 100}]
		)
		self.assertRaisesRegex(
			frappe.ValidationError,
			"has no default receivable account in INR",
			self.doc.create_opening_invoices,
			"Customer",
		)
		self.assertFalse(frappe.db.exists("Sales Invoice", {"company": self.company}))

	def test_sheet_names_match_an_account_only_when_sure(self):
		from erpnext.setup.doctype.company_setup.company_setup import (
			get_sheet_root_type,
			match_account_name,
			suggest_account,
		)

		def account(name, root_type, **values):
			return frappe._dict(
				{
					"name": f"{name} - X",
					"account_name": name,
					"root_type": root_type,
					"is_group": 0,
					"disabled": 0,
				}
				| values
			)

		accounts = [
			account("Furniture and Fixtures", "Asset"),
			account("Salary", "Expense"),
			account("Office Rent", "Expense"),
			account("Rent Deposit", "Asset"),
			account("Cash", "Asset"),
			account("Petty Cash", "Asset"),
		]
		expense = get_sheet_root_type(["Indirect Expenses"])
		self.assertEqual(expense, "Expense")
		# the same name, tidied: used
		self.assertEqual(
			match_account_name("Furniture & Fixtures A/c", accounts), "Furniture and Fixtures - X"
		)
		# close, not the same: only suggested, and only when one account fits
		self.assertIsNone(match_account_name("Salaries", accounts, expense))
		self.assertEqual(suggest_account("Salaries", accounts, expense), "Salary - X")
		# a name inside another is suggested only within the kind the sheet says: "Rent" may be the
		# rent paid or the deposit
		self.assertEqual(suggest_account("Rent", accounts, expense), "Office Rent - X")
		self.assertIsNone(suggest_account("Rent", accounts))
		self.assertIsNone(
			suggest_account("Cash", [*accounts, frappe._dict(account("Cash", "Asset"), name="Cash 2 - X")])
		)
		# an account of another kind is never matched, even by the same name
		self.assertIsNone(match_account_name("Salary", accounts, "Income"))

	def test_old_invoices_keep_their_dates_only_when_asked(self):
		from erpnext.setup.doctype.company_setup.opening_invoices import get_posting_dates

		# test sites have years up to 2050 only
		frappe.set_user("Administrator")
		self.doc.migration_date = "2074-01-01"
		row = {"party_name": self.customer.customer_name, "outstanding_amount": 100}
		self.doc.save_opening_invoices(
			"Customer",
			[row | {"invoice_number": "OB-OLD-1", "invoice_date": "2072-05-10", "due_date": "2072-06-09"}],
		)
		dates = get_posting_dates(self.doc, "Customer")
		self.assertEqual((dates["no_year"], dates["undated"], dates["kept"]), (1, 0, 0))

		self.doc.create_opening_invoices("Customer", keep_old_dates=1)
		self.assertEqual(
			[
				str(date)
				for date in frappe.db.get_value("Sales Invoice", "OB-OLD-1", ["posting_date", "due_date"])
			],
			["2072-05-10", "2072-06-09"],
		)
		# the years between it and those ERPNext had are added too, so none is missing
		self.assertTrue(get_fiscal_year("2061-05-10", company=self.company, boolean=True))

		# not asked: the day before migration, with the due date moved there
		self.doc.reload()
		self.doc.save_opening_invoices(
			"Customer",
			[row | {"invoice_number": "OB-OLD-2", "invoice_date": "2005-05-10", "due_date": "2005-06-09"}],
		)
		self.doc.create_opening_invoices("Customer")
		self.assertEqual(
			[
				str(date)
				for date in frappe.db.get_value("Sales Invoice", "OB-OLD-2", ["posting_date", "due_date"])
			],
			["2073-12-31", "2073-12-31"],
		)
		self.assertFalse(get_fiscal_year("2005-05-10", company=self.company, boolean=True))

	def allow_adding_customers(self):
		user = frappe.session.user
		frappe.set_user("Administrator")
		make_user(user, "Sales User")
		frappe.set_user(user)

	def test_a_new_customer_gets_the_group_picked_for_it(self):
		self.allow_adding_customers()
		group = frappe.get_all("Customer Group", filters={"is_group": 0}, pluck="name", limit=1)[0]
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": "Onboarding Grouped Customer",
					"outstanding_amount": 100,
					"create_party": 1,
					"party_group": group,
				}
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.assertEqual(
			frappe.db.get_value(
				"Customer", {"customer_name": "Onboarding Grouped Customer"}, "customer_group"
			),
			group,
		)

	def test_only_names_confirmed_as_new_are_added_once_and_complete(self):
		self.allow_adding_customers()
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{"party_name": "Onboarding Twice Customer", "outstanding_amount": 100, "create_party": 1},
				{"party_name": "Onboarding Twice Customer", "outstanding_amount": 200, "create_party": 1},
			],
		)
		self.doc.create_opening_invoices("Customer")
		# posting again from the same, older copy finds the rows posted
		self.assertRaises(frappe.ValidationError, self.doc.create_opening_invoices, "Customer")
		self.doc.reload()

		rows = self.doc.opening_sales_invoices
		self.assertEqual({row.status for row in rows}, {"Created"}, [row.error for row in rows])
		self.assertEqual(frappe.db.count("Customer", {"customer_name": "Onboarding Twice Customer"}), 1)
		self.assertEqual(rows[0].party, rows[1].party)

		self.doc.save_opening_invoices(
			"Customer",
			[
				{"party_name": "Onboarding Confirmed Customer", "outstanding_amount": 100, "create_party": 1},
				{"party_name": "Onboarding Unconfirmed Customer", "outstanding_amount": 200},
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		rows = {row.party_name: row for row in self.doc.opening_sales_invoices}
		self.assertEqual(
			(rows["Onboarding Confirmed Customer"].status, rows["Onboarding Unconfirmed Customer"].status),
			("Created", "Failed"),
		)
		self.assertFalse(frappe.db.exists("Customer", {"customer_name": "Onboarding Unconfirmed Customer"}))

		# like a custom app making a field mandatory, with no default to fill it
		make_property_setter("Customer", "tax_id", "reqd", 1, "Check")
		self.addCleanup(frappe.clear_cache, doctype="Customer")
		self.doc.save_opening_invoices(
			"Customer",
			[
				{"party_name": self.customer.customer_name, "outstanding_amount": 100},
				{
					"party_name": "Onboarding Incomplete Customer",
					"outstanding_amount": 200,
					"create_party": 1,
				},
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		rows = {row.party_name: row for row in self.doc.opening_sales_invoices}
		known, new = rows[self.customer.customer_name], rows["Onboarding Incomplete Customer"]
		self.assertEqual((known.status, new.status), ("Created", "Failed"))
		self.assertIn("Tax ID", new.error)
		self.assertFalse(frappe.db.exists("Customer", {"customer_name": "Onboarding Incomplete Customer"}))

	def test_a_name_confirmed_as_new_needs_permission_to_add_customers(self):
		frappe.set_user("Administrator")
		frappe.set_user(make_user("onboarding.no.customers@example.com", "Accounts Manager"))
		self.doc = frappe.get_doc("Company Setup", self.company)
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[{"party_name": "Onboarding Not Allowed Customer", "outstanding_amount": 100, "create_party": 1}],
		)
		self.assertFalse(self.doc.get_invoice_names("Customer")["can_create"])
		self.assertRaises(frappe.PermissionError, self.doc.create_opening_invoices, "Customer")

	def test_a_name_posted_on_an_earlier_sheet_is_linked_on_the_next(self):
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": "ONBOARDING DELHI SPELLING",
					"party": self.customer.name,
					"outstanding_amount": 100,
				}
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()
		self.assertEqual(self.doc.opening_sales_invoices[0].status, "Created")

		# the next office's sheet spells the name the same way
		self.doc.save_opening_invoices(
			"Customer", [{"party_name": "Onboarding Delhi Spelling", "outstanding_amount": 200}]
		)
		self.assertEqual(self.doc.opening_sales_invoices[-1].party, self.customer.name)
		self.assertEqual(self.doc.get_invoice_names("Customer")["new"], [])

	def test_sheet_for_new_customers_has_the_new_names_and_mandatory_columns(self):
		self.allow_adding_customers()
		make_property_setter("Customer", "tax_id", "reqd", 1, "Check")
		self.addCleanup(frappe.clear_cache, doctype="Customer")
		self.doc.save_opening_invoices(
			"Customer",
			[
				{"party_name": self.customer.customer_name, "outstanding_amount": 100},
				{"party_name": "Onboarding Sheet Customer", "outstanding_amount": 200},
				{"party_name": "Onboarding Sheet Customer", "outstanding_amount": 300},
			],
		)
		download_new_parties(self.company, "Customer")

		rows = list(load_workbook(BytesIO(frappe.response["filecontent"])).active.iter_rows(values_only=True))
		header = rows[0]
		self.assertEqual(header[:2], ("Customer Name", "Customer Type"))
		self.assertIn("Tax ID", header)
		self.assertIn("Mobile No (Contact)", header)
		self.assertEqual([row[0] for row in rows[1:]], ["Onboarding Sheet Customer"])

	def add_receivable_account(self, account_name):
		user = frappe.session.user
		frappe.set_user("Administrator")
		receivable = get_account(self.company, account_type="Receivable")
		account = frappe.get_doc(
			{
				"doctype": "Account",
				"account_name": account_name,
				"company": self.company,
				"parent_account": frappe.db.get_value("Account", receivable, "parent_account"),
				"account_type": "Receivable",
			}
		).insert()
		frappe.set_user(user)
		return receivable, account.name

	def test_invoices_are_set_against_the_trial_balance_by_account(self):
		from erpnext.setup.doctype.company_setup.opening_invoices import get_sheet_match

		_, export = self.add_receivable_account("Debtors - Export")
		default = frappe.get_cached_value("Company", self.company, "default_receivable_account")
		row = {"party_name": self.customer.customer_name, "party": self.customer.name}
		self.doc.save_opening_invoices(
			"Customer",
			[
				row | {"invoice_number": "OB-M-1", "outstanding_amount": 300},
				row | {"invoice_number": "OB-M-2", "outstanding_amount": 150, "account": export},
			],
		)

		def lines(match):
			return {line["account"]: (line["sheet"], line["entered"]) for line in match["lines"]}

		self.doc.sheet_accounts = {"Receivable": {default: 300, export: 200}}
		self.assertEqual(
			lines(get_sheet_match(self.doc, "Customer")), {default: (300, 300), export: (200, 150)}
		)

		# an account the trial balance doesn't have
		self.doc.sheet_accounts = {"Receivable": {export: 200}}
		self.assertEqual(lines(get_sheet_match(self.doc, "Customer"))[default], (None, 300))

		# customers listed by name: only the total can be set against the invoices
		self.doc.sheet_accounts = {"Receivable": {"": 500}}
		match = get_sheet_match(self.doc, "Customer")
		self.assertEqual(
			(match["by_account"], match["sheet"], match["entered"], match["lines"]), (False, 500, 450, [])
		)

	def test_posting_fixes_only_the_account_clashes_ticked(self):
		from erpnext.setup.doctype.company_setup.opening_invoices import get_account_clashes

		_, export = self.add_receivable_account("Debtors - Export")
		default = frappe.get_cached_value("Company", self.company, "default_receivable_account")
		self.doc.migration_date = "2026-04-01"
		row = {"party_name": self.customer.customer_name, "invoice_date": "2026-03-01"}
		self.doc.save_opening_invoices(
			"Customer", [row | {"invoice_number": "OB-C-1", "outstanding_amount": 300}], "a.xlsx"
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()
		self.doc.db_set("sheet_accounts", frappe.as_json({"Receivable": {export: 600}}))

		# a later sheet's ledger puts the customer on Export, where the trial balance is short
		self.doc.save_opening_invoices(
			"Customer",
			[row | {"invoice_number": "OB-C-2", "outstanding_amount": 200, "account": export}],
			"b.xlsx",
		)
		(clash,) = get_account_clashes(self.doc, "Customer")
		self.assertEqual(
			(clash["kind"], clash["account"], clash["from_account"], clash["said"], clash["parties"]),
			(
				"posted",
				export,
				default,
				True,
				[{"party": self.customer.name, "sources": ["a.xlsx"], "count": 1, "amount": 300}],
			),
		)
		# a fix the server did not find is ignored
		self.doc.create_opening_invoices(
			"Customer", [{"kind": "draft", "account": default, "from_account": export}]
		)
		self.doc.reload()
		self.assertEqual(frappe.db.get_value("Sales Invoice", "OB-C-1", "docstatus"), 1)

		self.doc.save_opening_invoices(
			"Customer",
			[row | {"invoice_number": "OB-C-3", "outstanding_amount": 100, "account": export}],
			"b.xlsx",
		)
		self.doc.create_opening_invoices(
			"Customer", [{"kind": "posted", "account": export, "from_account": default}]
		)
		self.doc.reload()
		# moved: deleted and posted again on Export under its own number
		self.assertEqual(
			{
				name: frappe.db.get_value("Sales Invoice", name, ["docstatus", "debit_to"])
				for name in ("OB-C-1", "OB-C-2", "OB-C-3")
			},
			{name: (1, export) for name in ("OB-C-1", "OB-C-2", "OB-C-3")},
		)
		self.assertEqual(self.doc.undone_sales_invoices, [])

		# now this sheet is the odd one: it says the default account, the earlier invoices are on Export
		self.doc.db_set("sheet_accounts", frappe.as_json({"Receivable": {export: 650}}))
		self.doc.save_opening_invoices(
			"Customer",
			[row | {"invoice_number": "OB-C-4", "outstanding_amount": 50, "account": default}],
			"c.xlsx",
		)
		(clash,) = get_account_clashes(self.doc, "Customer")
		self.assertEqual(
			(clash["kind"], clash["account"], clash["from_account"], clash["saved"], clash["earlier"]),
			("draft", export, default, False, True),
		)
		self.doc.create_opening_invoices(
			"Customer", [{"kind": "draft", "account": export, "from_account": default}]
		)
		self.assertEqual(frappe.db.get_value("Sales Invoice", "OB-C-4", "debit_to"), export)
		self.assertEqual(get_account_clashes(frappe.get_doc("Company Setup", self.company), "Customer"), [])

	def test_an_invoice_number_another_company_used_gets_this_companys_abbreviation(self):
		from erpnext.accounts.doctype.sales_invoice.test_sales_invoice import create_sales_invoice

		user = frappe.session.user
		frappe.set_user("Administrator")
		other = create_sales_invoice(company="_Test Company", do_not_save=1)
		other.insert(set_name="OB-SHARED-1")
		frappe.set_user(user)
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": self.customer.customer_name,
					"invoice_number": "OB-SHARED-1",
					"outstanding_amount": 100,
				}
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()
		abbr = frappe.get_cached_value("Company", self.company, "abbr")
		row = self.doc.opening_sales_invoices[0]
		self.assertEqual((row.status, row.invoice), ("Created", f"OB-SHARED-1-{abbr}"))

	def test_invoices_post_to_their_customers_account_and_keep_it(self):
		self.assertEqual(self.doc.get_invoice_names("Customer")["accounts"], [])
		receivable, export = self.add_receivable_account("Debtors - Export")
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": self.customer.customer_name,
					"account": export,
					"invoice_number": "OB-SINV-EXP-1",
					"outstanding_amount": 500,
				},
				{"party_name": self.customer.customer_name, "account": export, "outstanding_amount": -100},
			],
		)
		names = self.doc.get_invoice_names("Customer")
		self.assertEqual(set(names["accounts"]), {receivable, export})
		self.assertEqual(names["party_accounts"], {self.customer.name: receivable})
		self.assertEqual(names["default_account"], receivable)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		rows = self.doc.opening_sales_invoices
		self.assertEqual([row.status for row in rows], ["Created"] * 2, [row.error for row in rows])
		self.assertEqual(frappe.db.get_value("Sales Invoice", "OB-SINV-EXP-1", "debit_to"), export)
		advance = frappe.get_doc("Journal Entry", rows[1].invoice)
		self.assertEqual(next(line.account for line in advance.accounts if line.party), export)
		# posting leaves the customer's own account alone: only the user or an import sets it
		self.assertEqual(get_party_account("Customer", self.customer.name, self.company), receivable)

		# only the company's receivable or payable accounts can be picked
		income = get_account(self.company, root_type="Income")
		for account in (income, get_account(COMPANY, account_type="Receivable")):
			with self.assertRaises(frappe.ValidationError):
				self.doc.save_opening_invoices(
					"Customer",
					[
						{
							"party_name": self.customer.customer_name,
							"account": account,
							"outstanding_amount": 1,
						}
					],
				)
		# an employee advance is payable too, but an asset: never a supplier's account
		for account in get_invoice_accounts(self.company, "Supplier"):
			self.assertEqual(frappe.db.get_value("Account", account, "root_type"), "Liability")
		self.assertIn(export, get_invoice_accounts(self.company, "Customer"))

	def test_a_failed_invoice_is_undone_alone(self):
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": self.customer.customer_name,
					"invoice_number": "OB-DUP-1",
					"outstanding_amount": 100,
				},
				{
					"party_name": self.customer.customer_name,
					"invoice_number": "OB-DUP-1",
					"outstanding_amount": 200,
				},
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		first, second = self.doc.opening_sales_invoices
		self.assertEqual((first.status, second.status), ("Created", "Failed"))
		self.assertTrue(second.error)
		self.assertNotIn("<", second.error)
		self.doc.run_method("onload")
		step = self.doc.get_onload().steps[step_key(self.doc, "opening_sales_invoices")]
		self.assertEqual(
			step["summary"], "1 invoice posted · 1 invoice left to post: 1 failed, fix them and post again"
		)

	def test_posted_and_repeated_invoices_are_kept_once(self):
		result = self.doc.import_opening_invoices(
			"Customer",
			make_file(
				"twice.csv",
				[
					"Customer,Invoice No,Invoice Date,Outstanding",
					f"{self.customer.customer_name},OB-TWICE-1,2026-01-01,700",
					f"{self.customer.customer_name},OB-TWICE-1,2026-01-01,700",
				],
			),
		)
		self.assertEqual((result["added"], result["repeated"]), (1, 1))
		self.assertEqual(len(self.doc.opening_sales_invoices), 1)

		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{
					"party_name": self.customer.customer_name,
					"invoice_number": "OB-KEEP-1",
					"outstanding_amount": 100,
				}
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		# a client sends the form without the posted row
		self.doc.opening_sales_invoices = []
		self.doc.save()
		self.doc.reload()
		self.assertEqual(
			[(row.invoice_number, row.status) for row in self.doc.opening_sales_invoices],
			[("OB-KEEP-1", "Created")],
		)
		self.doc.migration_date = "2026-03-01"
		self.assertRaises(frappe.ValidationError, self.doc.save)
		self.doc.reload()

		# uploading a sheet again leaves out what it posted
		sheet = make_file(
			"receivables.csv",
			[
				"Customer,Invoice No,Invoice Date,Outstanding",
				f"{self.customer.customer_name},OB-AGAIN-1,2026-02-01,1200",
				"Onboarding New Customer,OB-AGAIN-2,2026-02-03,800",
			],
		)
		self.doc.import_opening_invoices("Customer", sheet)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()

		result = self.doc.import_opening_invoices("Customer", sheet)
		self.assertEqual(
			result,
			{
				"added": 1,
				"advances": 0,
				"skipped": 1,
				"repeated": 0,
				"unused_columns": [],
				"bad_dates": [],
				"unknown_ledgers": [],
				"post_to": None,
				"source": "receivables.csv",
			},
		)
		self.assertEqual(
			[(row.invoice_number, row.status) for row in self.doc.opening_sales_invoices],
			[("OB-KEEP-1", "Created"), ("OB-AGAIN-1", "Created"), ("OB-AGAIN-2", "Draft")],
		)

	def test_posting_adds_a_temporary_account_when_the_chart_has_none(self):
		# like an imported chart: no temporary account at all
		frappe.set_user("Administrator")
		frappe.delete_doc("Account", get_account(self.company, account_type="Temporary"))
		frappe.set_user("onboarding.accounts@example.com")
		self.doc.migration_date = "2026-04-01"
		self.doc.save()
		# invoices and balances each add it: undo the first so the second starts without one too
		frappe.db.savepoint("no_temporary_account")

		self.doc.save_opening_invoices(
			"Customer", [{"party_name": self.customer.customer_name, "outstanding_amount": 100}]
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.reload()
		self.assertEqual(
			self.doc.opening_sales_invoices[0].status, "Created", self.doc.opening_sales_invoices[0].error
		)
		frappe.db.rollback(save_point="no_temporary_account")

		self.doc.reload()
		self.doc.save_opening_balances(
			[{"account": get_account(self.company, account_type="Cash"), "debit": 1000}]
		)
		self.doc.post_opening_balances()
		account = frappe.db.get_value(
			"Account", {"company": self.company, "account_type": "Temporary", "disabled": 0}
		)
		self.assertEqual(company_setup_steps.get_temporary_opening_balance(self.company), -1000)
		self.assertEqual(frappe.db.get_value("Account", account, "root_type"), "Asset")

	def test_advances_and_supplier_bills_post_as_opening_entries(self):
		supplier = frappe.get_all("Supplier", pluck="supplier_name", limit=1)[0]
		self.doc.migration_date = "2026-04-01"
		self.doc.save_opening_invoices(
			"Customer",
			[
				{"party_name": self.customer.customer_name, "outstanding_amount": 1000},
				{"party_name": self.customer.customer_name, "outstanding_amount": -300},
			],
		)
		self.doc.save_opening_invoices(
			"Supplier",
			[
				{"party_name": supplier, "outstanding_amount": -200},
				{"party_name": supplier, "invoice_number": "BILL/77", "outstanding_amount": 500},
			],
		)
		self.doc.create_opening_invoices("Customer")
		self.doc.create_opening_invoices("Supplier")
		self.doc.reload()

		advances = (self.doc.opening_sales_invoices[1], self.doc.opening_purchase_invoices[0])
		self.assertEqual(
			[row.status for row in advances], ["Created", "Created"], [row.error for row in advances]
		)
		for row, debit, credit in ((advances[0], 0, 300), (advances[1], 200, 0)):
			je = frappe.get_doc("Journal Entry", row.invoice)
			self.assertEqual((je.voucher_type, je.is_opening, je.docstatus), ("Opening Entry", "Yes", 1))
			party_line = next(line for line in je.accounts if line.party)
			self.assertEqual(
				(party_line.party, party_line.is_advance, party_line.debit, party_line.credit),
				(row.party, "Yes", debit, credit),
			)

		# a supplier's bill number is kept as the bill's
		bill = self.doc.opening_purchase_invoices[1]
		self.assertEqual(bill.status, "Created", bill.error)
		self.assertEqual(frappe.db.get_value("Purchase Invoice", bill.invoice, "bill_no"), "BILL/77")

		# 1000 owed by a customer, less 300 paid in advance, plus 200 paid to a supplier, less 500 owed to it
		self.assertEqual(company_setup_steps.get_temporary_opening_balance(self.company), -400)
